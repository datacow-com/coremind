"""
Test suite for core.ingestion.nodes.parser.cpu_parser module.

Tests CPU-based text parsing including:
- PDF text and table extraction
- Office document parsing (DOCX, PPTX)
- Multi-page document handling
- Table Markdown conversion
- Concurrent processing control
"""

import io
import sys
from unittest.mock import Mock, patch, MagicMock
from typing import Any, Dict

import pytest

from core.ingestion.nodes.parser.cpu_parser import CpuTextParser


# Check if optional modules are available
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pptx
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


class TestCpuTextParser:
    """Test CPU text parser functionality."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for CPU parser tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "pdf",
            "raw_content": b"fake_pdf_content",
            "error_log": [],
            "progress": {},
        }


class TestPDFParsing:
    """Test PDF parsing functionality."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for PDF parser tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "pdf",
            "raw_content": b"fake_pdf_content",
            "error_log": [],
            "progress": {},
        }
    
    @pytest.mark.asyncio
    async def test_pdf_text_and_table_extraction(self, base_state):
        """TC-CP001: 验证PDF文本块和表格的正确提取 (P1)"""
        parser = CpuTextParser()
        
        # Mock PyMuPDF document with text blocks and tables
        mock_doc = MagicMock()
        mock_page = MagicMock()
        
        # Mock text blocks: (x0, y0, x1, y1, text, block_no, block_type)
        # block_type 0 = text
        mock_page.get_text.return_value = [
            (10, 10, 100, 30, "Header text", 0, 0),  # Text block
            (10, 40, 100, 80, "Body paragraph", 1, 0),  # Text block
        ]
        
        # Mock table detection
        mock_table = MagicMock()
        mock_table.extract.return_value = [
            ["Name", "Age", "City"],
            ["Alice", "25", "NYC"],
            ["Bob", "30", "LA"],
        ]
        mock_table.bbox = [10, 90, 100, 150]
        mock_page.find_tables.return_value = [mock_table]
        
        # Make doc iterable
        mock_doc.__iter__ = Mock(return_value=iter([mock_page]))
        mock_doc.__enter__ = Mock(return_value=mock_doc)
        mock_doc.__exit__ = Mock(return_value=None)
        
        with patch("core.ingestion.nodes.parser.cpu_parser.fitz.open", return_value=mock_doc):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        
        # Verify text blocks
        text_blocks = [b for b in blocks if b["type"] == "text"]
        assert len(text_blocks) == 2
        assert text_blocks[0]["content"] == "Header text"
        assert text_blocks[0]["page"] == 1
        assert text_blocks[0]["bbox"] == [10, 10, 100, 30]
        
        # Verify table blocks
        table_blocks = [b for b in blocks if b["type"] == "table"]
        assert len(table_blocks) == 1
        assert "Name | Age | City" in table_blocks[0]["content"]
        assert "Alice | 25 | NYC" in table_blocks[0]["content"]
        assert table_blocks[0]["page"] == 1
    
    @pytest.mark.asyncio
    async def test_table_markdown_conversion(self, base_state):
        """TC-CP002: 验证表格正确转换为Markdown格式 (P1)"""
        parser = CpuTextParser()
        
        # Test table conversion
        table_data = [
            ["Header 1", "Header 2", "Header 3"],
            ["Row 1 Col 1", "Row 1 Col 2", "Row 1 Col 3"],
            ["Row 2 Col 1", "Row 2 Col 2", "Row 2 Col 3"],
        ]
        
        markdown = parser._table_to_markdown(table_data)
        
        # Verify Markdown format
        lines = markdown.split("\n")
        assert "| Header 1 | Header 2 | Header 3 |" in lines[0]
        assert "| --- | --- | --- |" in lines[1]
        assert "| Row 1 Col 1 | Row 1 Col 2 | Row 1 Col 3 |" in lines[2]
        assert "| Row 2 Col 1 | Row 2 Col 2 | Row 2 Col 3 |" in lines[3]
    
    @pytest.mark.asyncio
    async def test_multi_page_page_numbering(self, base_state):
        """TC-CP003: 验证多页PDF的页码正确标记 (P1)"""
        parser = CpuTextParser()
        
        # Mock multi-page document
        mock_doc = MagicMock()
        
        # Create multiple pages
        pages = []
        for page_num in range(3):
            mock_page = MagicMock()
            mock_page.get_text.return_value = [
                (10, 10, 100, 30, f"Page {page_num + 1} content", 0, 0),
            ]
            mock_page.find_tables.return_value = []
            pages.append(mock_page)
        
        mock_doc.__iter__ = Mock(return_value=iter(pages))
        mock_doc.__enter__ = Mock(return_value=mock_doc)
        mock_doc.__exit__ = Mock(return_value=None)
        
        with patch("core.ingestion.nodes.parser.cpu_parser.fitz.open", return_value=mock_doc):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        
        # Verify page numbering
        page_nums = set(b["page"] for b in blocks)
        assert len(page_nums) == 3
        assert min(page_nums) == 1
        assert max(page_nums) == 3
        
        # Verify content per page
        page_1_blocks = [b for b in blocks if b["page"] == 1]
        assert "Page 1 content" in page_1_blocks[0]["content"]
    
    @pytest.mark.asyncio
    async def test_pdf_parsing_without_tables(self, base_state):
        """Test PDF parsing when find_tables is not available."""
        parser = CpuTextParser()
        
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_page.get_text.return_value = [
            (10, 10, 100, 30, "Text without tables", 0, 0),
        ]
        
        # Simulate older PyMuPDF without find_tables - raise AttributeError
        mock_page.find_tables.side_effect = AttributeError("find_tables not available")
        
        mock_doc.__iter__ = Mock(return_value=iter([mock_page]))
        mock_doc.__enter__ = Mock(return_value=mock_doc)
        mock_doc.__exit__ = Mock(return_value=None)
        
        with patch("core.ingestion.nodes.parser.cpu_parser.fitz.open", return_value=mock_doc):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        assert len(blocks) == 1
        assert blocks[0]["type"] == "text"
        assert blocks[0]["content"] == "Text without tables"


class TestOfficeDocumentParsing:
    """Test Office document parsing."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for Office parser tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "docx",
            "raw_content": b"fake_docx_content",
            "error_log": [],
            "progress": {},
        }
    
    @pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
    @pytest.mark.asyncio
    async def test_docx_table_and_style_extraction(self, base_state):
        """TC-CP005: 验证DOCX文档的表格和样式信息 (P1)"""
        base_state["file_type"] = "docx"
        parser = CpuTextParser()
        
        # Mock python-docx Document
        mock_doc = MagicMock()
        
        # Mock paragraphs with different styles
        mock_para1 = MagicMock()
        mock_para1.text = "Document Title"
        mock_para1.style.name = "Heading 1"
        
        mock_para2 = MagicMock()
        mock_para2.text = "This is a paragraph"
        mock_para2.style.name = "Normal"
        
        mock_doc.paragraphs = [mock_para1, mock_para2]
        
        # Mock table
        mock_table = MagicMock()
        mock_row1 = MagicMock()
        mock_row1.cells = [MagicMock(text="Col1"), MagicMock(text="Col2")]
        mock_row2 = MagicMock()
        mock_row2.cells = [MagicMock(text="Data1"), MagicMock(text="Data2")]
        
        mock_table.rows = [mock_row1, mock_row2]
        mock_doc.tables = [mock_table]
        
        with patch("core.ingestion.nodes.parser.cpu_parser.Document", return_value=mock_doc):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        
        # Verify headers
        headers = [b for b in blocks if b["type"] == "header"]
        assert len(headers) == 1
        assert headers[0]["content"] == "Document Title"
        assert headers[0]["style"] == "Heading 1"
        
        # Verify tables
        tables = [b for b in blocks if b["type"] == "table"]
        assert len(tables) == 1
        table_content = tables[0]["content"]
        assert "Col1 | Col2" in table_content
        assert "Data1 | Data2" in table_content
    
    @pytest.mark.skipif(not HAS_PPTX, reason="python-pptx not installed")
    @pytest.mark.asyncio
    async def test_pptx_slide_parsing(self, base_state):
        """TC-CP006: 验证PPTX每页内容正确提取 (P1)"""
        base_state["file_type"] = "pptx"
        parser = CpuTextParser()
        
        # Mock python-pptx Presentation
        mock_prs = MagicMock()
        
        # Create mock slides
        slides = []
        for slide_num in range(3):
            mock_slide = MagicMock()
            
            # Mock shapes with text
            mock_shape = MagicMock()
            mock_shape.text = f"Slide {slide_num + 1} title"
            mock_shape.has_table = False
            
            mock_slide.shapes = [mock_shape]
            slides.append(mock_slide)
        
        mock_prs.slides = slides
        
        with patch("core.ingestion.nodes.parser.cpu_parser.Presentation", return_value=mock_prs):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        
        # Verify slide pages
        pages = set(b["page"] for b in blocks)
        assert len(pages) == 3
        assert min(pages) == 1
        assert max(pages) == 3
        
        # Verify slide content
        slide_1_blocks = [b for b in blocks if b["page"] == 1]
        assert "Slide 1 title" in slide_1_blocks[0]["content"]
    
    @pytest.mark.skipif(not HAS_PPTX, reason="python-pptx not installed")
    @pytest.mark.asyncio
    async def test_pptx_table_in_slides(self, base_state):
        """Test PPTX table extraction from slides."""
        base_state["file_type"] = "pptx"
        parser = CpuTextParser()
        
        mock_prs = MagicMock()
        mock_slide = MagicMock()
        
        # Mock shape with table
        mock_shape = MagicMock()
        mock_shape.has_table = True
        mock_shape.text = ""
        
        # Mock table in shape
        mock_table = MagicMock()
        mock_row = MagicMock()
        mock_row.cells = [MagicMock(text="Header1"), MagicMock(text="Header2")]
        mock_table.rows = [mock_row]
        mock_shape.table = mock_table
        
        mock_slide.shapes = [mock_shape]
        mock_prs.slides = [mock_slide]
        
        with patch("core.ingestion.nodes.parser.cpu_parser.Presentation", return_value=mock_prs):
            result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        table_blocks = [b for b in blocks if b["type"] == "table"]
        assert len(table_blocks) == 1
        assert "Header1 | Header2" in table_blocks[0]["content"]


class TestOtherFormats:
    """Test parsing of other document formats."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for other format tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "txt",
            "raw_content": b"test content",
            "error_log": [],
            "progress": {},
        }
    
    @pytest.mark.asyncio
    async def test_markdown_parsing(self, base_state):
        """Test Markdown file parsing."""
        base_state["file_type"] = "md"
        base_state["raw_content"] = b"# Header\n\nParagraph 1\n\nParagraph 2"
        
        parser = CpuTextParser()
        result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        assert len(blocks) == 3  # Header + 2 paragraphs
        assert blocks[0]["content"] == "# Header"
        assert blocks[1]["content"] == "Paragraph 1"
        assert blocks[2]["content"] == "Paragraph 2"
    
    @pytest.mark.asyncio
    async def test_html_parsing(self, base_state):
        """Test HTML file parsing."""
        base_state["file_type"] = "html"
        base_state["raw_content"] = b"<h1>Title</h1><p>Paragraph</p><script>alert('test')</script>"
        
        parser = CpuTextParser()
        result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        # Should extract text and ignore script
        text_content = " ".join(b["content"] for b in blocks)
        assert "Title" in text_content
        assert "Paragraph" in text_content
        assert "alert" not in text_content  # Script should be ignored
    
    @pytest.mark.asyncio
    async def test_email_parsing(self, base_state):
        """Test email (.eml) file parsing."""
        base_state["file_type"] = "eml"
        email_content = b"""From: sender@example.com
To: recipient@example.com
Subject: Test Email

This is the email body content.
"""
        base_state["raw_content"] = email_content
        
        parser = CpuTextParser()
        result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        assert len(blocks) == 1
        content = blocks[0]["content"]
        assert "Subject: Test Email" in content
        assert "From: sender@example.com" in content
        assert "This is the email body content." in content
    
    @pytest.mark.asyncio
    async def test_plain_text_parsing(self, base_state):
        """Test plain text file parsing."""
        base_state["file_type"] = "txt"
        base_state["raw_content"] = b"Line 1\n\nLine 2\n\nLine 3"
        
        parser = CpuTextParser()
        result = await parser(base_state)
        
        blocks = result["parsed_blocks"]
        assert len(blocks) == 3
        assert blocks[0]["content"] == "Line 1"
        assert blocks[1]["content"] == "Line 2"
        assert blocks[2]["content"] == "Line 3"


class TestErrorHandling:
    """Test error handling and edge cases."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for error handling tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "pdf",
            "raw_content": b"fake_pdf_content",
            "error_log": [],
            "progress": {},
        }
    
    @pytest.mark.asyncio
    async def test_pdf_parsing_error(self, base_state):
        """Test PDF parsing error handling."""
        parser = CpuTextParser()
        
        with patch("core.ingestion.nodes.parser.cpu_parser.fitz.open", side_effect=Exception("PDF parsing error")):
            result = await parser(base_state)
        
        # Should record error and return empty blocks
        assert len(result["error_log"]) == 1
        assert "PDF parsing error" in result["error_log"][0]["error"]
        assert result["parsed_blocks"] == []
    
    @pytest.mark.asyncio
    async def test_docx_import_error(self, base_state):
        """Test DOCX parsing when python-docx is not available."""
        base_state["file_type"] = "docx"
        parser = CpuTextParser()
        
        # Mock the import inside _parse_docx to raise ImportError
        with patch.dict(sys.modules, {'docx': None}):
            with patch.object(parser, '_parse_docx', side_effect=ImportError("python-docx is required")):
                result = await parser(base_state)
        
        # Should record import error
        assert len(result["error_log"]) == 1
        assert "python-docx is required" in result["error_log"][0]["error"]
    
    @pytest.mark.asyncio
    async def test_pptx_import_error(self, base_state):
        """Test PPTX parsing when python-pptx is not available."""
        base_state["file_type"] = "pptx"
        parser = CpuTextParser()
        
        # Mock the import inside _parse_pptx to raise ImportError
        with patch.dict(sys.modules, {'pptx': None}):
            with patch.object(parser, '_parse_pptx', side_effect=ImportError("python-pptx is required")):
                result = await parser(base_state)
        
        # Should record import error
        assert len(result["error_log"]) == 1
        assert "python-pptx is required" in result["error_log"][0]["error"]
    
    @pytest.mark.asyncio
    async def test_unsupported_file_type(self, base_state):
        """Test handling of unsupported file types."""
        base_state["file_type"] = "unknown"
        parser = CpuTextParser()
        
        result = await parser(base_state)
        
        # Should return empty blocks for unsupported types
        assert result["parsed_blocks"] == []
    
    @pytest.mark.asyncio
    async def test_empty_content_handling(self, base_state):
        """Test handling of empty content."""
        base_state["raw_content"] = None
        parser = CpuTextParser()
        
        result = await parser(base_state)
        
        # Should return original state when no content
        assert result == base_state
    
    @pytest.mark.asyncio
    async def test_corrupted_docx_handling(self, base_state):
        """Test handling of corrupted DOCX files."""
        base_state["file_type"] = "docx"
        parser = CpuTextParser()
        
        # Mock _parse_docx to raise exception
        with patch.object(parser, '_parse_docx', side_effect=Exception("Corrupted file")):
            result = await parser(base_state)
        
        assert len(result["error_log"]) == 1
        assert "Corrupted file" in result["error_log"][0]["error"]


class TestConcurrencyControl:
    """Test concurrent processing control."""
    
    @pytest.fixture
    def base_state(self) -> Dict[str, Any]:
        """Base state for concurrency tests."""
        return {
            "channel_id": "test_channel",
            "task_id": "task_001",
            "file_type": "pdf",
            "raw_content": b"fake_pdf_content",
            "error_log": [],
            "progress": {},
        }
    
    @pytest.mark.asyncio
    async def test_concurrent_parsing_limit(self, base_state):
        """TC-CP004: 验证CPU解析器的并发控制 (P2)"""
        import asyncio
        
        parser = CpuTextParser()
        
        # Track concurrent executions
        concurrent_count = 0
        max_concurrent = 0
        lock = asyncio.Lock()
        
        original_parse_pdf = parser._parse_pdf
        
        def mock_parse_pdf(content):
            nonlocal concurrent_count, max_concurrent
            concurrent_count += 1
            max_concurrent = max(max_concurrent, concurrent_count)
            result = [{"type": "text", "content": "test", "page": 1, "bbox": None}]
            concurrent_count -= 1
            return result
        
        # Create multiple parsing tasks
        states = []
        for i in range(5):
            state = base_state.copy()
            state["task_id"] = f"task_{i}"
            state["error_log"] = []
            states.append(state)
        
        with patch.object(parser, '_parse_pdf', side_effect=mock_parse_pdf):
            # Process all states concurrently
            tasks = [parser(state) for state in states]
            results = await asyncio.gather(*tasks)
        
        # Verify all tasks completed
        assert len(results) == 5
        # Verify parsing happened (max_concurrent may be 1 if sync execution)
        for result in results:
            assert "parsed_blocks" in result


class TestTableConversion:
    """Test table conversion utilities."""
    
    def test_empty_table_conversion(self):
        """Test conversion of empty table."""
        parser = CpuTextParser()
        
        result = parser._table_to_markdown([])
        assert result == ""
    
    def test_single_row_table_conversion(self):
        """Test conversion of single row table."""
        parser = CpuTextParser()
        
        table_data = [["Header1", "Header2", "Header3"]]
        result = parser._table_to_markdown(table_data)
        
        lines = result.split("\n")
        assert len(lines) == 2  # Header + separator
        assert "| Header1 | Header2 | Header3 |" in lines[0]
        assert "| --- | --- | --- |" in lines[1]
    
    def test_table_with_empty_cells(self):
        """Test conversion of table with empty cells."""
        parser = CpuTextParser()
        
        table_data = [
            ["Name", "Age", "City"],
            ["Alice", "", "NYC"],
            ["", "30", ""],
        ]
        result = parser._table_to_markdown(table_data)
        
        lines = result.split("\n")
        assert "| Alice |  | NYC |" in lines[2]
        assert "|  | 30 |  |" in lines[3]


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
